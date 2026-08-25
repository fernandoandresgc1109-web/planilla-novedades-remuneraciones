import csv
from datetime import date
from decimal import Decimal
from io import BytesIO, StringIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from .models import (
    AFP,
    Colaborador,
    Exportacion,
    InstitucionSalud,
    Novedad,
    PeriodoLiquidacion,
    Sucursal,
    TipoNovedad,
)


class ExportacionNovedadesTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.usuario = get_user_model().objects.create_user(
            username="usuario_exportacion",
            password="ClaveSegura123!",
        )

        cls.sucursal = Sucursal.objects.create(
            codigo="SUC-EXPORT",
            nombre="Sucursal Exportación Ficticia",
            direccion="Dirección ficticia",
            activo=True,
        )

        cls.afp = AFP.objects.create(
            codigo="AFP-EXPORT",
            nombre="AFP Exportación Ficticia",
            activo=True,
        )

        cls.salud = InstitucionSalud.objects.create(
            codigo="SALUD-EXPORT",
            nombre="Salud Exportación Ficticia",
            tipo=InstitucionSalud.Tipo.FONASA,
            activo=True,
        )

        cls.colaborador = Colaborador.objects.create(
            rut="88.888.888-8",
            nombres="Elena",
            apellidos="Exportación",
            sucursal=cls.sucursal,
            afp=cls.afp,
            institucion_salud=cls.salud,
            cargas_familiares=0,
            seguro_cesantia=True,
            activo=True,
        )

        cls.periodo = PeriodoLiquidacion.objects.create(
            sucursal=cls.sucursal,
            anio=2026,
            mes=8,
            fecha_inicio=date(2026, 8, 1),
            estado=PeriodoLiquidacion.Estado.ABIERTO,
        )

        cls.tipo_bono = TipoNovedad.objects.create(
            codigo="BONO-EXPORT",
            nombre="Bono de exportación ficticio",
            unidad_medida=TipoNovedad.UnidadMedida.PESOS,
            naturaleza=TipoNovedad.Naturaleza.HABER,
            activo=True,
        )

        momento_validacion = timezone.now()

        cls.novedad_validada = Novedad.objects.create(
            periodo=cls.periodo,
            colaborador=cls.colaborador,
            tipo_novedad=cls.tipo_bono,
            monto=Decimal("50000.00"),
            observacion="Registro validado para exportación.",
            estado=Novedad.Estado.VALIDADA,
            creado_por=cls.usuario,
            validado_por=cls.usuario,
            validado_en=momento_validacion,
        )

        cls.novedad_anulada = Novedad.objects.create(
            periodo=cls.periodo,
            colaborador=cls.colaborador,
            tipo_novedad=cls.tipo_bono,
            monto=Decimal("10000.00"),
            observacion=(
                "Registro anulado que no debe exportarse."
            ),
            estado=Novedad.Estado.ANULADA,
            creado_por=cls.usuario,
            validado_por=cls.usuario,
            validado_en=momento_validacion,
            anulado_por=cls.usuario,
            anulado_en=timezone.now(),
            motivo_anulacion=(
                "Registro ficticio anulado para la prueba."
            ),
        )

    def setUp(self):
        self.client.force_login(self.usuario)
        self.url = reverse("novedades:exportar_novedades")

    def exportar(self, formato):
        return self.client.post(
            self.url,
            {
                "periodo": self.periodo.pk,
                "formato": formato,
            },
        )

    def test_ruta_requiere_autenticacion(self):
        self.client.logout()

        respuesta = self.client.get(self.url)

        self.assertRedirects(
            respuesta,
            f"{reverse('login')}?next={self.url}",
        )

    def test_pagina_muestra_formulario_y_periodo_disponible(self):
        respuesta = self.client.get(self.url)

        self.assertEqual(respuesta.status_code, 200)
        self.assertTemplateUsed(
            respuesta,
            "novedades/exportar_novedades.html",
        )
        self.assertContains(respuesta, "Exportar novedades")
        self.assertContains(respuesta, "Excel")
        self.assertContains(respuesta, "CSV")

        formulario = respuesta.context["formulario"]

        self.assertIn(
            self.periodo,
            formulario.fields["periodo"].queryset,
        )

    def test_periodo_sin_novedades_validadas_no_aparece(self):
        Novedad.objects.filter(
            pk=self.novedad_validada.pk
        ).update(
            estado=Novedad.Estado.ANULADA,
            anulado_por=self.usuario,
            anulado_en=timezone.now(),
            motivo_anulacion=(
                "Registro anulado para comprobar el formulario."
            ),
        )

        respuesta = self.client.get(self.url)
        formulario = respuesta.context["formulario"]

        self.assertNotIn(
            self.periodo,
            formulario.fields["periodo"].queryset,
        )

    def test_exportacion_se_bloquea_si_existen_borradores(self):
        Novedad.objects.create(
            periodo=self.periodo,
            colaborador=self.colaborador,
            tipo_novedad=self.tipo_bono,
            monto=Decimal("25000.00"),
            observacion="Borrador que debe bloquear la exportación.",
            estado=Novedad.Estado.BORRADOR,
            creado_por=self.usuario,
        )

        respuesta = self.exportar(Exportacion.Formato.CSV)

        self.assertEqual(respuesta.status_code, 200)
        self.assertIn(
            "periodo",
            respuesta.context["formulario"].errors,
        )
        self.assertContains(
            respuesta,
            "todavía contiene novedades en borrador",
        )
        self.assertEqual(Exportacion.objects.count(), 0)
        self.assertNotIn(
            "Content-Disposition",
            respuesta.headers,
        )

        self.periodo.refresh_from_db()

        self.assertEqual(
            self.periodo.estado,
            PeriodoLiquidacion.Estado.ABIERTO,
        )
        self.assertIsNone(self.periodo.fecha_cierre)

    def test_csv_incluye_solo_validadas_y_registra_auditoria(self):
        respuesta = self.exportar(Exportacion.Formato.CSV)

        self.assertEqual(respuesta.status_code, 200)
        self.assertTrue(
            respuesta["Content-Type"].startswith("text/csv")
        )
        self.assertEqual(
            respuesta["Content-Disposition"],
            (
                'attachment; filename="'
                'novedades_suc-export_2026_08.csv"'
            ),
        )
        self.assertTrue(
            respuesta.content.startswith(b"\xef\xbb\xbf")
        )

        contenido = respuesta.content.decode("utf-8-sig")
        filas = list(
            csv.reader(
                StringIO(contenido),
                delimiter=";",
            )
        )

        self.assertEqual(len(filas), 2)
        self.assertEqual(filas[0][0], "Sucursal")
        self.assertEqual(filas[0][-1], "Validado en")
        self.assertEqual(filas[1][3], "88.888.888-8")
        self.assertEqual(filas[1][6], "BONO-EXPORT")
        self.assertEqual(filas[1][13], "50000.00")
        self.assertEqual(filas[1][15], "Validada")
        self.assertNotIn(
            "Registro anulado que no debe exportarse",
            contenido,
        )

        exportacion = Exportacion.objects.get()

        self.assertEqual(
            exportacion.formato,
            Exportacion.Formato.CSV,
        )
        self.assertEqual(exportacion.cantidad_registros, 1)
        self.assertEqual(exportacion.generado_por, self.usuario)

        self.periodo.refresh_from_db()

        self.assertEqual(
            self.periodo.estado,
            PeriodoLiquidacion.Estado.EXPORTADO,
        )
        self.assertIsNotNone(self.periodo.fecha_cierre)

    def test_xlsx_contiene_datos_y_formato_esperado(self):
        respuesta = self.exportar(Exportacion.Formato.XLSX)

        self.assertEqual(respuesta.status_code, 200)
        self.assertTrue(
            respuesta.content.startswith(b"PK")
        )
        self.assertIn(
            "novedades_suc-export_2026_08.xlsx",
            respuesta["Content-Disposition"],
        )

        libro = load_workbook(BytesIO(respuesta.content))
        hoja = libro["Novedades"]

        self.assertEqual(hoja.max_row, 2)
        self.assertEqual(hoja.max_column, 18)
        self.assertEqual(hoja["A1"].value, "Sucursal")
        self.assertEqual(hoja["R1"].value, "Validado en")
        self.assertTrue(hoja["A1"].font.bold)
        self.assertEqual(hoja.freeze_panes, "A2")
        self.assertEqual(hoja["G2"].value, "BONO-EXPORT")
        self.assertEqual(hoja["N2"].value, 50000)
        self.assertEqual(
            hoja["O2"].value,
            "Registro validado para exportación.",
        )

        valores = [
            celda.value
            for fila in hoja.iter_rows()
            for celda in fila
        ]

        self.assertNotIn(
            "Registro anulado que no debe exportarse.",
            valores,
        )

    def test_reexportacion_conserva_fecha_de_cierre(self):
        primera_respuesta = self.exportar(
            Exportacion.Formato.CSV
        )

        self.assertEqual(
            primera_respuesta.status_code,
            200,
        )

        self.periodo.refresh_from_db()
        primera_fecha_cierre = self.periodo.fecha_cierre

        segunda_respuesta = self.exportar(
            Exportacion.Formato.XLSX
        )

        self.assertEqual(
            segunda_respuesta.status_code,
            200,
        )

        self.periodo.refresh_from_db()

        self.assertEqual(Exportacion.objects.count(), 2)
        self.assertEqual(
            self.periodo.fecha_cierre,
            primera_fecha_cierre,
        )
        self.assertEqual(
            self.periodo.estado,
            PeriodoLiquidacion.Estado.EXPORTADO,
        )

    def test_textos_peligrosos_no_se_exportan_como_formulas(self):
        novedad = Novedad.objects.get(
            pk=self.novedad_validada.pk
        )
        novedad.observacion = "=2+2"
        novedad.save(update_fields=["observacion"])

        respuesta_csv = self.exportar(
            Exportacion.Formato.CSV
        )
        contenido_csv = respuesta_csv.content.decode(
            "utf-8-sig"
        )

        self.assertIn("'=2+2", contenido_csv)

        respuesta_xlsx = self.exportar(
            Exportacion.Formato.XLSX
        )
        libro = load_workbook(
            BytesIO(respuesta_xlsx.content)
        )
        hoja = libro["Novedades"]

        self.assertEqual(hoja["O2"].value, "'=2+2")