import csv
from decimal import Decimal
from io import BytesIO, StringIO

from django.utils import timezone
from django.utils.text import slugify
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


COLUMNAS = (
    "Sucursal",
    "Año",
    "Mes",
    "RUT",
    "Nombres",
    "Apellidos",
    "Código de novedad",
    "Tipo de novedad",
    "Unidad",
    "Naturaleza",
    "Fecha de inicio",
    "Fecha de término",
    "Cantidad",
    "Monto",
    "Observación",
    "Estado",
    "Validado por",
    "Validado en",
)

TIPO_CSV = "text/csv; charset=utf-8"
TIPO_XLSX = (
    "application/vnd.openxmlformats-officedocument."
    "spreadsheetml.sheet"
)


def texto_seguro(valor):
    """
    Evita que Excel interprete textos como fórmulas.
    """

    if valor is None:
        return ""

    texto = str(valor)

    if texto.startswith(("=", "+", "-", "@")):
        return f"'{texto}"

    return texto


def fecha_texto(valor):
    if valor is None:
        return ""

    return valor.strftime("%d/%m/%Y")


def fecha_hora_texto(valor):
    if valor is None:
        return ""

    return timezone.localtime(valor).strftime("%d/%m/%Y %H:%M")


def preparar_fila(novedad):
    usuario_validador = ""

    if novedad.validado_por:
        usuario_validador = novedad.validado_por.username

    return [
        texto_seguro(novedad.periodo.sucursal.nombre),
        novedad.periodo.anio,
        novedad.periodo.mes,
        texto_seguro(novedad.colaborador.rut),
        texto_seguro(novedad.colaborador.nombres),
        texto_seguro(novedad.colaborador.apellidos),
        texto_seguro(novedad.tipo_novedad.codigo),
        texto_seguro(novedad.tipo_novedad.nombre),
        texto_seguro(
            novedad.tipo_novedad.get_unidad_medida_display()
        ),
        texto_seguro(
            novedad.tipo_novedad.get_naturaleza_display()
        ),
        fecha_texto(novedad.fecha_inicio),
        fecha_texto(novedad.fecha_termino),
        novedad.cantidad if novedad.cantidad is not None else "",
        novedad.monto if novedad.monto is not None else "",
        texto_seguro(novedad.observacion),
        texto_seguro(novedad.get_estado_display()),
        texto_seguro(usuario_validador),
        fecha_hora_texto(novedad.validado_en),
    ]


def generar_csv(novedades):
    contenido = StringIO(newline="")

    escritor = csv.writer(
        contenido,
        delimiter=";",
        quoting=csv.QUOTE_MINIMAL,
        lineterminator="\n",
    )

    escritor.writerow(COLUMNAS)

    for novedad in novedades:
        escritor.writerow(preparar_fila(novedad))

    return contenido.getvalue().encode("utf-8-sig")


def generar_xlsx(novedades):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Novedades"

    libro.properties.creator = "Planilla de Novedades"
    libro.properties.title = "Exportación de novedades"

    hoja.append(COLUMNAS)

    for novedad in novedades:
        fila = preparar_fila(novedad)

        fila_excel = [
            float(valor) if isinstance(valor, Decimal) else valor
            for valor in fila
        ]

        hoja.append(fila_excel)

    color_encabezado = PatternFill(
        fill_type="solid",
        fgColor="047857",
    )
    fuente_encabezado = Font(
        color="FFFFFF",
        bold=True,
    )

    for celda in hoja[1]:
        celda.fill = color_encabezado
        celda.font = fuente_encabezado
        celda.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    hoja.row_dimensions[1].height = 24

    for fila in hoja.iter_rows(
        min_row=2,
        min_col=13,
        max_col=13,
    ):
        fila[0].number_format = "0.00"

    for fila in hoja.iter_rows(
        min_row=2,
        min_col=14,
        max_col=14,
    ):
        fila[0].number_format = '#,##0.00'

    for columna in hoja.columns:
        longitud_maxima = 0
        letra = get_column_letter(columna[0].column)

        for celda in columna:
            valor = "" if celda.value is None else str(celda.value)
            longitud_maxima = max(longitud_maxima, len(valor))

        hoja.column_dimensions[letra].width = min(
            longitud_maxima + 2,
            45,
        )

    salida = BytesIO()
    libro.save(salida)
    salida.seek(0)

    return salida.getvalue()


def construir_nombre_archivo(periodo, formato):
    codigo_sucursal = slugify(periodo.sucursal.codigo)

    if not codigo_sucursal:
        codigo_sucursal = f"sucursal-{periodo.sucursal_id}"

    extension = formato.lower()

    return (
        f"novedades_{codigo_sucursal}_"
        f"{periodo.anio}_{periodo.mes:02d}.{extension}"
    )


def generar_archivo(formato, novedades):
    if formato == "CSV":
        return generar_csv(novedades), TIPO_CSV

    if formato == "XLSX":
        return generar_xlsx(novedades), TIPO_XLSX

    raise ValueError("El formato solicitado no está permitido.")